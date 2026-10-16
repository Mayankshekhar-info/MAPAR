"""
Automated Unit and Integration Tests for Uploads Engine.
Validates Template CRUD, Column Schemas, Versioning, Dynamic Excel Generation,
Authoritative Validation, Hierarchy Checking, Circle Scopes, and Transactional Commit.
"""

import io
from typing import Any
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook  # type: ignore

from apps.imports.models import (
    UploadTemplate,
    UploadTemplateVersion,
    UploadTemplateColumn,
    ImportBatch,
)
from apps.imports.services.template_service import TemplateService
from apps.imports.services.excel_generation_service import ExcelGenerationService
from apps.imports.services.upload_validation_service import UploadValidationService
from apps.imports.services.import_execution_service import ImportExecutionService
from apps.organization.models import MainCompany, SubCompany, Project, Circle, Domain, Department

User = get_user_model()


class UploadsEngineTestCase(TestCase):
    def setUp(self):
        # Create organization hierarchy
        self.company = MainCompany.objects.create(name="Apex Enterprise", code="APEX")
        self.sub_company = SubCompany.objects.create(main_company=self.company, name="Apex Energy", code="APEX_ENG")
        self.project = Project.objects.create(sub_company=self.sub_company, name="Solar Field Alpha", code="PRJ_SOLAR")
        self.circle_north = Circle.objects.create(project=self.project, name="North Circle", code="CIRC_NORTH")
        self.circle_south = Circle.objects.create(project=self.project, name="South Circle", code="CIRC_SOUTH")
        self.domain = Domain.objects.create(circle=self.circle_north, name="Operations", code="OPS")

        # Create user assigned to North Circle
        user_model: Any = get_user_model()
        self.user = user_model.objects.create_user(
            username="analyst_test",
            email="analyst@mapar.test",
            password="testpassword123",
            role="ANALYST",
            circle=self.circle_north,
        )

    def test_create_template_with_dynamic_columns(self):
        """Test template creation with dynamic columns and initial version v1."""
        template_data = {
            "name": "Operational KPI Roster",
            "description": "Monthly operational metrics ingestion template",
            "data_level": "CIRCLE",
            "target_entity": "KPI_VALUES",
            "columns": [
                {
                    "name": "circle_code",
                    "display_label": "Circle Code",
                    "data_type": "REFERENCE",
                    "is_required": True,
                    "validation_rules": {"reference_level": "CIRCLE"},
                },
                {
                    "name": "metric_name",
                    "display_label": "Metric Name",
                    "data_type": "TEXT",
                    "is_required": True,
                },
                {
                    "name": "target_value",
                    "display_label": "Target Value",
                    "data_type": "DECIMAL",
                    "is_required": False,
                },
            ],
        }

        template: Any = TemplateService.create_template(template_data, user=self.user)
        self.assertIsNotNone(template.id)
        self.assertEqual(template.name, "Operational KPI Roster")
        self.assertEqual(template.versions.count(), 1)

        v1: Any = template.versions.first()
        self.assertEqual(v1.version_number, 1)
        self.assertEqual(v1.columns.count(), 3)

        # Check column order
        cols = list(v1.columns.all().order_by("column_order"))
        self.assertEqual(cols[0].display_label, "Circle Code")
        self.assertEqual(cols[1].display_label, "Metric Name")
        self.assertEqual(cols[2].display_label, "Target Value")

    def test_dynamic_excel_generation(self):
        """Verify dynamic .xlsx generation directly reflects database column schema."""
        template: Any = TemplateService.create_template({
            "name": "Staff Count Import",
            "data_level": "CIRCLE",
            "columns": [
                {"name": "employee_id", "display_label": "Employee ID", "data_type": "TEXT", "is_required": True, "is_unique": True},
                {"name": "headcount", "display_label": "Headcount", "data_type": "INTEGER", "is_required": True},
            ],
        }, user=self.user)

        version = template.versions.first()
        excel_buffer = ExcelGenerationService.generate_template_workbook(version)
        self.assertGreater(excel_buffer.getbuffer().nbytes, 1000)

        # Inspect generated workbook
        wb = load_workbook(filename=excel_buffer, data_only=True)
        self.assertIn("Data Entry", wb.sheetnames)
        self.assertIn("Column Guide & Instructions", wb.sheetnames)

        ws = wb["Data Entry"]
        # Required headers get asterisk suffix
        self.assertEqual(ws.cell(row=1, column=1).value, "Employee ID *")
        self.assertEqual(ws.cell(row=1, column=2).value, "Headcount *")

    def test_upload_and_authoritative_validation(self):
        """Test upload validation pipeline: headers, required fields, unique constraints, and organization circle reference."""
        template: Any = TemplateService.create_template({
            "name": "Energy Output Import",
            "data_level": "CIRCLE",
            "columns": [
                {"name": "circle_code", "display_label": "Circle Code", "data_type": "REFERENCE", "is_required": True, "validation_rules": {"reference_level": "CIRCLE"}},
                {"name": "batch_tag", "display_label": "Batch Tag", "data_type": "TEXT", "is_required": True, "is_unique": True},
                {"name": "megawatts", "display_label": "Megawatts", "data_type": "DECIMAL", "is_required": True},
            ],
        }, user=self.user)

        version = template.versions.first()

        # Build an Excel file in-memory
        wb = Workbook()
        ws: Any = wb.active
        ws.title = "Data Entry"
        ws.append(["Circle Code *", "Batch Tag *", "Megawatts *"])
        # Row 2: Valid
        ws.append(["CIRC_NORTH", "TAG-101", 450.5])
        # Row 3: Invalid circle code
        ws.append(["INVALID_CIRCLE_999", "TAG-102", 320.0])
        # Row 4: Duplicate batch tag & missing megawatts
        ws.append(["CIRC_NORTH", "TAG-101", None])

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        uploaded = SimpleUploadedFile("energy_data.xlsx", excel_buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        batch: Any = UploadValidationService.process_and_validate_upload(
            version=version,
            uploaded_file=uploaded,
            user=self.user,
        )

        self.assertEqual(batch.total_rows, 3)
        self.assertEqual(batch.valid_rows, 1)
        self.assertEqual(batch.error_rows, 2)
        self.assertGreater(batch.validation_errors.count(), 0)

        # Inspect specific error records
        errors = list(batch.validation_errors.all())
        error_msgs = [e.error_message for e in errors]
        self.assertTrue(any("does not exist" in m for m in error_msgs))
        self.assertTrue(any("Duplicate value" in m for m in error_msgs))
        self.assertTrue(any("is required" in m for m in error_msgs))

        # Test error report generation
        err_buffer = ImportExecutionService.generate_error_report(batch)
        self.assertGreater(err_buffer.getbuffer().nbytes, 1000)

    def test_template_versioning(self):
        """Test creating new version v2 without destroying historical v1."""
        template: Any = TemplateService.create_template({
            "name": "Equipment Inspection",
            "columns": [{"name": "asset_code", "display_label": "Asset Code", "data_type": "TEXT", "is_required": True}],
        }, user=self.user)

        v1: Any = template.versions.first()
        self.assertEqual(v1.version_number, 1)

        v2: Any = TemplateService.create_new_version(
            template_id=template.id,
            data={
                "changelog": "Added inspection status and hours",
                "columns": [
                    {"name": "asset_code", "display_label": "Asset Code", "data_type": "TEXT", "is_required": True},
                    {"name": "inspection_hours", "display_label": "Inspection Hours", "data_type": "INTEGER", "is_required": False},
                ],
            },
            user=self.user,
        )

        self.assertEqual(v2.version_number, 2)
        self.assertEqual(v1.columns.count(), 1)
        self.assertEqual(v2.columns.count(), 2)
        self.assertEqual(template.versions.count(), 2)

    def test_uploads_to_cockpit_and_slides_integration(self):
        """Verify dynamic template -> upload -> commit -> KPIValue -> Cockpit & Slide View query flow."""
        import datetime
        from django.db import connection
        from apps.dashboard.models import KPI, KPIValue, ReportingPeriod
        from apps.dashboard.services.cockpit_service import CockpitService
        from apps.dashboard.services.slide_table_service import SlideTableService

        # Ensure unmanaged tables exist in test database
        with connection.cursor() as cursor:
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS organization_reportingperiod (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name VARCHAR(100) UNIQUE,
                    period_type VARCHAR(20),
                    start_date DATE,
                    end_date DATE,
                    freeze_at DATETIME,
                    is_frozen BOOLEAN DEFAULT 0,
                    is_current BOOLEAN DEFAULT 0,
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP
                );
            """)
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS kpis_kpi (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    code VARCHAR(20) UNIQUE,
                    name VARCHAR(200),
                    definition TEXT DEFAULT '',
                    calculation_formula TEXT DEFAULT '',
                    unit VARCHAR(50) DEFAULT '',
                    direction VARCHAR(20) DEFAULT 'HIGHER_BETTER',
                    frequency VARCHAR(20) DEFAULT 'MONTHLY',
                    risk_category VARCHAR(100) DEFAULT '',
                    priority VARCHAR(20) DEFAULT 'MEDIUM',
                    default_owner VARCHAR(100) DEFAULT '',
                    is_executive BOOLEAN DEFAULT 0,
                    is_active BOOLEAN DEFAULT 1,
                    display_order INTEGER DEFAULT 0,
                    domain_id INTEGER,
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
                );
            """)
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS kpis_kpivalue (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    kpi_id INTEGER,
                    circle_id INTEGER,
                    period_id INTEGER,
                    actual_value NUMERIC(15, 2),
                    target_value NUMERIC(15, 2),
                    previous_value NUMERIC(15, 2),
                    numerator_value NUMERIC(15, 2),
                    denominator_value NUMERIC(15, 2),
                    variance_to_target NUMERIC(15, 2),
                    variance_percentage NUMERIC(8, 2),
                    rag_status VARCHAR(20) DEFAULT 'GREEN',
                    trend VARCHAR(20) DEFAULT 'FLAT',
                    notes TEXT DEFAULT '',
                    entered_by_id INTEGER,
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    updated_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE (kpi_id, circle_id, period_id)
                );
            """)

        period = ReportingPeriod.objects.create(
            name="September 2026",
            period_type="MONTHLY",
            start_date=datetime.date(2026, 9, 1),
            end_date=datetime.date(2026, 9, 30),
            is_current=True,
        )

        template: Any = TemplateService.create_template({
            "name": "Manpower Ingestion Roster",
            "data_level": "CIRCLE",
            "columns": [
                {"name": "circle_code", "display_label": "Circle Code", "data_type": "REFERENCE", "is_required": True, "validation_rules": {"reference_level": "CIRCLE"}},
                {"name": "metric_code", "display_label": "Resource Code", "data_type": "TEXT", "is_required": True},
                {"name": "required", "display_label": "Required Headcount", "data_type": "INTEGER", "is_required": True},
                {"name": "available", "display_label": "Available Headcount", "data_type": "INTEGER", "is_required": True},
            ],
        }, user=self.user)

        version = template.versions.first()

        # Build workbook with valid records
        wb = Workbook()
        ws: Any = wb.active
        ws.title = "Data Entry"
        ws.append(["Circle Code *", "Resource Code *", "Required Headcount *", "Available Headcount *"])
        ws.append(["CIRC_NORTH", "TECH_ENG_01", 100, 85])
        ws.append(["CIRC_NORTH", "TECH_ENG_02", 50, 50])

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        uploaded = SimpleUploadedFile(
            "manpower_upload.xlsx",
            excel_buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        batch: Any = UploadValidationService.process_and_validate_upload(
            version=version,
            uploaded_file=uploaded,
            reporting_period_id=period.id,
            user=self.user,
        )

        self.assertEqual(batch.valid_rows, 2)
        self.assertEqual(batch.error_rows, 0)

        # Commit batch into MAPAR live data layer
        commit_res = ImportExecutionService.commit_batch(batch, user=self.user)
        self.assertTrue(commit_res["success"])
        self.assertEqual(commit_res["committed_rows"], 2)

        # 1. Verify Management Cockpit consumes newly committed records
        summary = CockpitService.get_master_summary(user=self.user, circle_id=self.circle_north.id, period_id=period.id)
        self.assertGreater(summary["metrics_summary"]["total_kpis"], 0)
        self.assertGreater(len(summary["rows"]), 0)

        # 2. Verify Slide View consumes newly committed records
        slide_table = SlideTableService.get_master_table(user=self.user, circle_id=self.circle_north.id, period_id=period.id)
        self.assertGreater(slide_table["grand_total"]["record_count"], 0)
        self.assertEqual(slide_table["grand_total"]["required"], 150.0)
        self.assertEqual(slide_table["grand_total"]["available"], 135.0)
        self.assertEqual(slide_table["grand_total"]["gap"], 15.0)

        # 3. Verify Master -> Detail click navigation
        data_rows = [r for r in slide_table["rows"] if r.get("row_type") == "data"]
        self.assertGreater(len(data_rows), 0)
        first_rec_id = data_rows[0]["id"]
        detail = SlideTableService.get_record_detail(record_id=first_rec_id, user=self.user)
        self.assertIsNotNone(detail)
        self.assertEqual(detail["hierarchy"]["circle_code"], "CIRC_NORTH")

    def test_upload_with_datetime_and_decimal_cells(self):
        """Test uploading an Excel file containing Python datetime and Decimal objects."""
        import datetime
        from decimal import Decimal

        template: Any = TemplateService.create_template({
            "name": "Site Ops Log",
            "data_level": "CIRCLE",
            "columns": [
                {"name": "site_id", "display_label": "Site ID", "data_type": "TEXT"},
                {"name": "commission_date", "display_label": "Commission Date", "data_type": "TEXT"},
                {"name": "capacity", "display_label": "Capacity", "data_type": "DECIMAL"},
            ],
        }, user=self.user)

        version = template.versions.first()

        wb = Workbook()
        ws: Any = wb.active
        ws.title = "Data Entry"
        ws.append(["Site ID", "Commission Date", "Capacity"])
        # Append Python datetime and Decimal objects directly (as openpyxl does for Excel dates)
        ws.append(["SITE-001", datetime.datetime(2026, 10, 15, 10, 30), Decimal("125.75")])
        ws.append(["SITE-002", datetime.date(2026, 10, 14), 85.5])

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        uploaded = SimpleUploadedFile(
            "site_ops.xlsx",
            excel_buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        batch: Any = UploadValidationService.process_and_validate_upload(
            version=version,
            uploaded_file=uploaded,
            user=self.user,
        )

        self.assertEqual(batch.total_rows, 2)
        self.assertEqual(batch.valid_rows, 2)
        staged_rows = list(batch.staged_rows.all().order_by("row_index"))
        self.assertEqual(len(staged_rows), 2)
        self.assertEqual(staged_rows[0].raw_data["commission_date"], "2026-10-15 10:30:00")
        self.assertEqual(float(staged_rows[0].raw_data["capacity"]), 125.75)
        self.assertEqual(staged_rows[1].raw_data["commission_date"], "2026-10-14")


