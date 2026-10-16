"""
Tests for Upload Excel Template with Exact Specified Columns Only.
Verifies character-for-character column headers, single worksheet, zero sample rows,
zero dummy data, zero instruction sheets, and full compatibility with Uploads module.
"""

import io
from openpyxl import Workbook, load_workbook
from django.test import TestCase
from rest_framework.test import APIClient
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.imports.models import (
    UploadTemplate,
    UploadTemplateVersion,
    UploadTemplateColumn,
)
from apps.imports.services.excel_generation_service import ExcelGenerationService
from apps.imports.services.upload_validation_service import UploadValidationService

User = get_user_model()

REQUIRED_COLUMNS = [
    "SL",
    "Entity",
    "Site Type",
    "Circle Name",
    "BSNL Site ID",
    "Site Name/Village Name",
    "BA",
    "OA",
    "District",
    "Site Commissioning /Hoto Status",
    "O&M Status",
    "EB Status",
    "DG Status",
    "Site On Air Status",
    "IE Name",
    "CI Name",
    "O&M Name",
    "CBH Name",
    "Outage min",
    "Uptime",
]


class UploadExcelTemplateTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="template_tester",
            email="tester@mapar.internal",
            password="testpassword123",
            role="SUPER_ADMIN"
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_database_template_columns_count_and_order(self):
        """Verifies that the template in the Uploads module has exactly 20 columns in exact order."""
        template = UploadTemplate.objects.filter(code="UPLOAD_EXCEL_TEMPLATE").first()
        self.assertIsNotNone(template, "UPLOAD_EXCEL_TEMPLATE should exist in database.")

        version = template.versions.filter(is_active=True).first()
        self.assertIsNotNone(version, "Template should have an active version.")

        db_columns = list(version.columns.order_by("column_order", "id"))
        self.assertEqual(len(db_columns), 20, "Template must have exactly 20 columns.")

        for idx, (expected_header, col) in enumerate(zip(REQUIRED_COLUMNS, db_columns), start=1):
            self.assertEqual(col.column_order, idx, f"Column order mismatch for index {idx}")
            self.assertEqual(col.display_label, expected_header, f"Header mismatch at position {idx}")

    def test_excel_generation_clean_structure(self):
        """Verifies that the generated Excel contains exactly 1 sheet, 20 columns, max_row 1, and no extra sheets/sample rows."""
        template = UploadTemplate.objects.get(code="UPLOAD_EXCEL_TEMPLATE")
        version = template.versions.filter(is_active=True).first()

        buffer = ExcelGenerationService.generate_template_workbook(version)
        self.assertIsNotNone(buffer)

        wb = load_workbook(filename=buffer, data_only=True)

        # 1. Exactly 1 sheet (no extra sheets)
        self.assertEqual(len(wb.sheetnames), 1, "Workbook must contain exactly 1 worksheet.")
        self.assertNotIn("Column Guide & Instructions", wb.sheetnames)

        ws = wb.active
        # 2. Exactly 1 header row (no sample rows, no dummy data)
        self.assertEqual(ws.max_row, 1, "Template must only have header row (max_row == 1), no sample rows.")

        # 3. Exactly 20 columns
        self.assertEqual(ws.max_column, 20, "Template must have exactly 20 columns.")

        # 4. Character-for-character check of every header
        generated_headers = [cell.value for cell in ws[1]]
        self.assertEqual(len(generated_headers), 20)
        for idx, (expected, actual) in enumerate(zip(REQUIRED_COLUMNS, generated_headers), start=1):
            self.assertEqual(
                actual,
                expected,
                f"Column {idx} mismatch: expected '{expected}', got '{actual}'"
            )

    def test_template_download_api_view(self):
        """Verifies downloading the template via GET /api/v1/uploads/templates/<id>/download/."""
        template = UploadTemplate.objects.get(code="UPLOAD_EXCEL_TEMPLATE")
        url = f"/api/v1/uploads/templates/{template.id}/download/"

        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        wb = load_workbook(filename=io.BytesIO(response.content), data_only=True)
        self.assertEqual(len(wb.sheetnames), 1)
        ws = wb.active
        self.assertEqual(ws.max_row, 1)
        self.assertEqual(ws.max_column, 20)

        headers = [c.value for c in ws[1]]
        self.assertEqual(headers, REQUIRED_COLUMNS)

    def test_upload_compatibility_with_uploads_module(self):
        """Verifies that an Excel file created from this template is compatible with the Uploads module pipeline."""
        template = UploadTemplate.objects.get(code="UPLOAD_EXCEL_TEMPLATE")
        version = template.versions.filter(is_active=True).first()

        # Create workbook with the 20 headers and 1 sample row
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(REQUIRED_COLUMNS)
        ws.append([
            1, "Enterprise", "Macro", "DELHI", "DEL001",
            "Connaught Place", "BA-North", "OA-Central", "New Delhi",
            "Commissioned", "Active", "EB Live", "DG Standby",
            "On Air", "Ramesh", "Suresh", "Mahesh", "CBH-01",
            0.5, 99.8
        ])

        file_buf = io.BytesIO()
        wb.save(file_buf)
        file_buf.seek(0)

        uploaded_file = SimpleUploadedFile(
            name="site_operations_test.xlsx",
            content=file_buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        batch = UploadValidationService.process_and_validate_upload(
            version=version,
            uploaded_file=uploaded_file,
            user=self.user
        )

        self.assertIsNotNone(batch)
        self.assertEqual(batch.total_rows, 1)
        self.assertEqual(batch.staged_rows.count(), 1)

    def test_update_existing_version_without_creating_new_version(self):
        """Verifies editing an existing version updates columns in-place without creating a new version."""
        from apps.imports.services.template_service import TemplateService

        template = UploadTemplate.objects.create(
            name="Version Edit Test Template",
            code="VER_EDIT_TEST",
            data_level="CIRCLE"
        )
        version1 = UploadTemplateVersion.objects.create(
            template=template,
            version_number=1,
            changelog="Original v1"
        )
        UploadTemplateColumn.objects.create(
            version=version1,
            name="col_a",
            display_label="Col A",
            data_type="TEXT",
            column_order=0
        )

        self.assertEqual(template.versions.count(), 1)

        # Edit version 1 in place
        update_data = {
            "changelog": "Edited in place v1",
            "columns": [
                {"name": "col_a", "display_label": "Col A Updated", "data_type": "TEXT", "column_order": 0},
                {"name": "col_b", "display_label": "Col B New", "data_type": "DECIMAL", "column_order": 1},
            ]
        }
        updated = TemplateService.update_version(template.id, version1.id, update_data, user=self.user)

        # Must still have only 1 version (no new version created)
        self.assertEqual(template.versions.count(), 1)
        self.assertEqual(updated.id, version1.id)
        self.assertEqual(updated.version_number, 1)
        self.assertEqual(updated.changelog, "Edited in place v1")
        self.assertEqual(updated.columns.count(), 2)

        # Test API endpoint
        res = self.client.put(
            f"/api/v1/uploads/templates/{template.id}/versions/{version1.id}/",
            {
                "changelog": "API Edited",
                "columns": [
                    {"display_label": "Col A", "data_type": "TEXT", "column_order": 0}
                ]
            },
            format="json"
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(template.versions.count(), 1)
        self.assertEqual(res.data["version"]["changelog"], "API Edited")

    def test_delete_old_version(self):
        """Verifies deleting an old uploaded version removes only that version without affecting other versions."""
        from apps.imports.services.template_service import TemplateService
        from django.core.exceptions import ValidationError

        template = UploadTemplate.objects.create(
            name="Version Delete Test Template",
            code="VER_DEL_TEST",
            data_level="CIRCLE"
        )
        v1 = UploadTemplateVersion.objects.create(
            template=template,
            version_number=1,
            changelog="v1"
        )
        v2 = UploadTemplateVersion.objects.create(
            template=template,
            version_number=2,
            changelog="v2",
            is_active=True
        )

        self.assertEqual(template.versions.count(), 2)

        # Delete v1 (old version)
        result = TemplateService.delete_version(template.id, v1.id, user=self.user)
        self.assertTrue(result["success"])

        # Check template versions
        self.assertEqual(template.versions.count(), 1)
        remaining = template.versions.first()
        self.assertEqual(remaining.id, v2.id)
        self.assertEqual(remaining.version_number, 2)

        # Cannot delete the only remaining version
        with self.assertRaises(ValidationError):
            TemplateService.delete_version(template.id, v2.id, user=self.user)

        # Test API delete endpoint
        # Create another version so we have 2 again
        v3 = UploadTemplateVersion.objects.create(
            template=template,
            version_number=3,
            changelog="v3"
        )
        res = self.client.delete(f"/api/v1/uploads/templates/{template.id}/versions/{v2.id}/")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(template.versions.count(), 1)
        self.assertEqual(template.versions.first().version_number, 3)

