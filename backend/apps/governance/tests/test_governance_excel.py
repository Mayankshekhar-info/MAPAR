import io
import pytest
from datetime import timedelta
from django.utils import timezone
import openpyxl
from apps.organization.models import MainCompany, SubCompany, Project, Circle
from apps.identity.models import User
from apps.governance.models import GovernanceActionPoint
from apps.governance.services.governance_excel_service import GovernanceExcelService


@pytest.fixture
def excel_setup(db):
    company = MainCompany.objects.create(name="Apex Group", code="APEX")
    sub_co = SubCompany.objects.create(main_company=company, name="Apex Solar", code="ASOL")
    project = Project.objects.create(sub_company=sub_co, name="Solar Project Alpha", code="SPA")
    circle1 = Circle.objects.create(project=project, name="Rajasthan Circle", code="RJ", region="North")
    circle2 = Circle.objects.create(project=project, name="Maharashtra Circle", code="MH", region="West")

    admin = User.objects.create_superuser(
        username="admin_xl",
        email="admin_xl@mapar.com",
        password="password123",
        role=User.Role.SUPER_ADMIN,
    )
    user1 = User.objects.create_user(
        username="mgupta",
        email="mgupta@mapar.com",
        password="password123",
        role=User.Role.ACTION_OWNER,
        circle=circle1,
        first_name="Mayank",
        last_name="Gupta",
    )
    return {
        "admin": admin,
        "user1": user1,
        "circle1": circle1,
        "circle2": circle2,
    }


@pytest.mark.django_db
def test_excel_import_success_and_validation(excel_setup):
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Action Tracker"

    headers = [
        "Action ID", "Circle", "Function/Area", "Action Point", "Key Requirement",
        "Owner", "Raised Date", "Target Closure Date", "Status", "Priority",
        "Actual Closure Date", "Last Update"
    ]
    ws.append(headers)

    today = timezone.localdate()
    target = today + timedelta(days=10)

    # Valid row
    ws.append([
        "",
        "RJ",
        "Operations",
        "Deploy inverter cooling fan kits",
        "Thermal alarms resolved",
        "mgupta",
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        "Open",
        "High",
        "",
        "PO issued to supplier",
    ])

    # Save to memory
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is True, res.get("errors")
    assert res["created_count"] == 1
    assert res["errors"] == []

    action = GovernanceActionPoint.objects.get(action_point="Deploy inverter cooling fan kits")
    assert action.circle.code == "RJ"
    assert action.owner is not None
    assert action.owner.username == "mgupta"
    assert action.priority == "High"
    assert action.status == "Open"
    assert action.ageing_days == 0
    assert action.is_escalated is False


@pytest.mark.django_db
def test_excel_import_catches_invalid_circle_and_date(excel_setup):
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Action Tracker"

    headers = [
        "Action ID", "Circle", "Function/Area", "Action Point", "Key Requirement",
        "Owner", "Raised Date", "Target Closure Date", "Status", "Priority",
        "Actual Closure Date", "Last Update"
    ]
    ws.append(headers)

    # Invalid Circle and Invalid Target Date
    ws.append([
        "",
        "NON_EXISTENT_CIRCLE_XYZ",
        "Operations",
        "Some action description",
        "",
        "mgupta",
        "2026-10-01",
        "invalid-date-format",
        "Open",
        "High",
        "",
        "",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is False
    assert len(res["errors"]) >= 2
    fields_with_errors = [e["field"] for e in res["errors"]]
    assert "Circle" in fields_with_errors
    assert "Target Closure Date" in fields_with_errors


@pytest.mark.django_db
def test_export_canonical_workbook_structure_and_formulas(excel_setup):
    """
    Verifies that export_governance_workbook generates the canonical multi-sheet workbook:
    - Open Points Tracker
    - Owner_wise_Dasboard (verifying formula references Column F - Owner, not Column B - Circle)
    - State-wise Dashboard
    - Ageing Summary
    - Target Revision History
    - How to Use
    """
    # Create action items
    GovernanceActionPoint.objects.create(
        action_id="ACT-00001",
        circle=excel_setup["circle1"],
        owner=excel_setup["user1"],
        action_point="Thermal calibration",
        target_closure_date=timezone.localdate() + timedelta(days=10),
        current_target_date=timezone.localdate() + timedelta(days=10),
        status="Open",
        priority="High",
        created_by=excel_setup["admin"],
    )

    stream = GovernanceExcelService.export_governance_workbook(excel_setup["admin"])
    wb = openpyxl.load_workbook(stream, data_only=False)

    expected_sheets = [
        "Open Points Tracker",
        "Owner_wise_Dasboard",
        "State-wise Dashboard",
        "Ageing Summary",
        "Target Revision History",
        "How to Use",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Expected sheet '{s}' missing from workbook."

    # Verify Owner_wise_Dasboard references Column F (Owner), NOT Column B (Circle)!
    ws_owner = wb["Owner_wise_Dasboard"]
    formula_total = str(ws_owner.cell(row=4, column=3).value)
    assert "$F$4:$F$" in formula_total, f"Owner-wise formula must reference column F (Owner), got: {formula_total}"
    assert "$B$4:$B$" not in formula_total, "Owner-wise formula must NOT reference column B (Circle)!"

    # Verify State-wise Dashboard references Column B (Circle)
    ws_state = wb["State-wise Dashboard"]
    formula_state = str(ws_state.cell(row=4, column=3).value)
    assert "$B$4:$B$" in formula_state, f"State-wise formula must reference column B (Circle), got: {formula_state}"


@pytest.mark.django_db
def test_excel_import_blocks_closure_without_accepted_evidence(excel_setup):
    """
    Verifies that closing a task via Excel import is rejected if mandatory evidence is unaccepted.
    """
    action = GovernanceActionPoint.objects.create(
        action_id="ACT-00099",
        circle=excel_setup["circle1"],
        owner=excel_setup["user1"],
        action_point="Inverter audit",
        target_closure_date=timezone.localdate() + timedelta(days=5),
        current_target_date=timezone.localdate() + timedelta(days=5),
        status="Open",
        evidence_required=True,
        evidence_status=GovernanceActionPoint.EvidenceStatus.NOT_SUBMITTED,
        created_by=excel_setup["admin"],
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Open Points Tracker"

    headers = [
        "S.No", "Circle", "Function/Area", "Action Point", "Key Requirement / Deliverable",
        "Owner", "Raised Date", "Target Closure Date", "Current Target", "Revision Count",
        "Status", "Actual Closure Date", "Ageing (Days)", "Ageing Bucket", "Priority",
        "Escalate to Business Head?", "Weeks in Review", "Last Update"
    ]
    ws.append(headers)
    ws.append([
        "ACT-00099",
        "RJ",
        "Operations",
        "Inverter audit",
        "Complete thermal scan",
        "mgupta",
        "2026-10-01",
        "2026-10-15",
        "2026-10-15",
        0,
        "Closed",  # Attempting to mark closed!
        "2026-10-20",
        0,
        "0-7 Days",
        "High",
        "N",
        0,
        "Done",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    # Action owner attempts closure without accepted evidence
    res = GovernanceExcelService.import_excel(stream, excel_setup["user1"])
    assert res["success"] is False
    assert any("evidence is required and has not been accepted" in str(e["message"]) for e in res["errors"])


@pytest.mark.django_db
def test_download_template_only_one_sheet():
    """
    Requirement 1: Verify Download Template produces an Excel file containing ONLY
    the 'Open Points Tracker' sheet, with existing columns and order unchanged.
    """
    stream = GovernanceExcelService.generate_template()
    wb = openpyxl.load_workbook(stream)

    assert wb.sheetnames == ["Open Points Tracker"], f"Template must contain only 'Open Points Tracker', found: {wb.sheetnames}"

    ws = wb["Open Points Tracker"]
    # Check headers on row 3
    headers = [cell.value for cell in ws[3]]
    assert headers == GovernanceExcelService.TRACKER_HEADERS, "Header columns and order must match canonical TRACKER_HEADERS"

    # Verify no dashboard, history, or reference sheets exist in the template
    excluded_sheets = [
        "Owner_wise_Dasboard",
        "State-wise Dashboard",
        "Ageing Summary",
        "Target Revision History",
        "How to Use",
        "Reference Data",
    ]
    for s in excluded_sheets:
        assert s not in wb.sheetnames, f"Sheet '{s}' must NOT be in the downloaded template!"


@pytest.mark.django_db
def test_excel_import_assigns_owner_by_username(excel_setup):
    """
    Requirement 2 & 3: In Open Points Tracker, Owner accepts username of existing user.
    Resolves to canonical database user ID and links task.
    Task appears in Audit & Action Owners (view_mode='my_tasks') for the assigned user.
    """
    from apps.governance.selectors.governance_selectors import get_governance_actions

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Open Points Tracker"
    ws.append(GovernanceExcelService.TRACKER_HEADERS)

    today = timezone.localdate()
    target = today + timedelta(days=14)

    ws.append([
        "ACT-00501",
        "RJ",
        "Operations",
        "Upgrade field inverter firmware",
        "All firmware upgraded to v2.4",
        "mgupta",  # Username
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        0,
        "Open",
        "",
        0,
        "0-7 Days",
        "High",
        "N",
        0,
        "Firmware patches downloaded",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is True, res.get("errors")
    assert res["created_count"] == 1

    action = GovernanceActionPoint.objects.get(action_id="ACT-00501")
    assert action.owner == excel_setup["user1"]
    assert action.owner is not None
    assert action.owner.id == excel_setup["user1"].id
    assert action.owner_name_override == ""
    assert action.circle == excel_setup["circle1"]
    assert action.priority == "High"
    assert action.status == "Open"

    # Verify task appears in Audit & Action Owners for the assigned user (view_mode='my_tasks')
    my_tasks = list(get_governance_actions({"view_mode": "my_tasks"}, excel_setup["user1"]))
    assert len(my_tasks) == 1
    assert my_tasks[0].action_id == "ACT-00501"
    assert my_tasks[0].owner == excel_setup["user1"]
    assert my_tasks[0].action_point == "Upgrade field inverter firmware"


@pytest.mark.django_db
def test_excel_import_assigns_owner_by_email(excel_setup):
    """
    Requirement 2: In Open Points Tracker, Owner accepts email address of existing user.
    Resolves to canonical database user ID and links task.
    """
    from apps.governance.selectors.governance_selectors import get_governance_actions

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Open Points Tracker"
    ws.append(GovernanceExcelService.TRACKER_HEADERS)

    today = timezone.localdate()
    target = today + timedelta(days=21)

    ws.append([
        "ACT-00502",
        "RJ",
        "Finance",
        "Verify subcontractor billing reconciliation",
        "Signed reconciliation statements",
        "mgupta@mapar.com",  # Email address
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        0,
        "WIP",
        "",
        0,
        "0-7 Days",
        "Medium",
        "N",
        0,
        "Initial invoices verified",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is True, res.get("errors")
    assert res["created_count"] == 1

    action = GovernanceActionPoint.objects.get(action_id="ACT-00502")
    assert action.owner == excel_setup["user1"]
    assert action.owner is not None
    assert action.owner.id == excel_setup["user1"].id
    assert action.owner_name_override == ""

    # Appears in assigned user's Audit & Action Owners list
    my_tasks = list(get_governance_actions({"view_mode": "my_tasks"}, excel_setup["user1"]))
    assert any(t.action_id == "ACT-00502" for t in my_tasks)


@pytest.mark.django_db
def test_excel_import_rejects_unmatched_owner_with_row_error(excel_setup):
    """
    Requirement 2 & 6: If the owner cannot be matched to a unique existing user,
    show a clear row-level validation error instead of assigning the task incorrectly.
    No task created.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Open Points Tracker"
    ws.append(GovernanceExcelService.TRACKER_HEADERS)

    today = timezone.localdate()
    target = today + timedelta(days=14)

    ws.append([
        "ACT-00503",
        "RJ",
        "Operations",
        "Test unmatched user error handling",
        "Deliverable 1",
        "nonexistent_user_999@xyz.com",  # Unmatched owner
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        0,
        "Open",
        "",
        0,
        "0-7 Days",
        "High",
        "N",
        0,
        "",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is False
    assert len(res["errors"]) >= 1
    owner_error = next((e for e in res["errors"] if e["field"] == "Owner"), None)
    assert owner_error is not None
    assert "nonexistent_user_999@xyz.com" in owner_error["message"]
    assert "does not match any existing MAPAR user" in owner_error["message"]

    # Confirm no task was created
    assert not GovernanceActionPoint.objects.filter(action_id="ACT-00503").exists()


@pytest.mark.django_db
def test_excel_import_rejects_empty_owner_with_row_error(excel_setup):
    """
    Requirement 2: Empty owner reports row-level error.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Open Points Tracker"
    ws.append(GovernanceExcelService.TRACKER_HEADERS)

    today = timezone.localdate()
    target = today + timedelta(days=14)

    ws.append([
        "ACT-00504",
        "RJ",
        "Operations",
        "Missing owner task",
        "Deliverable",
        "",  # Empty owner
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        0,
        "Open",
        "",
        0,
        "0-7 Days",
        "Low",
        "N",
        0,
        "",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is False
    owner_error = next((e for e in res["errors"] if e["field"] == "Owner"), None)
    assert owner_error is not None
    assert "Owner is required" in owner_error["message"]
    assert not GovernanceActionPoint.objects.filter(action_id="ACT-00504").exists()


@pytest.mark.django_db
def test_assigned_user_portal_update_synchronizes_with_governance(excel_setup):
    """
    Requirements 4 & 5:
    Assigned user updates task progress through the portal (e.g., weekly update).
    Saved changes persist, maintain audit trail, and reflect in Governance and Action.
    """
    from apps.governance.services.governance_weekly_update_service import GovernanceWeeklyUpdateService
    from apps.governance.selectors.governance_selectors import get_governance_actions

    # 1. Create task via Excel import assigned to user1 (mgupta)
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Open Points Tracker"
    ws.append(GovernanceExcelService.TRACKER_HEADERS)

    today = timezone.localdate()
    target = today + timedelta(days=20)

    ws.append([
        "ACT-00505",
        "RJ",
        "Operations",
        "Install perimeter security cameras",
        "CCTV coverage active across all gates",
        "mgupta",
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        0,
        "Open",
        "",
        0,
        "0-7 Days",
        "High",
        "N",
        0,
        "PO issued",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is True

    action = GovernanceActionPoint.objects.get(action_id="ACT-00505")
    assert action.current_progress_percent == 0
    assert action.status == "Open"

    # 2. Assigned user logs progress update through the portal workflow
    update_obj = GovernanceWeeklyUpdateService.submit_weekly_update(
        action_point=action,
        data={
            "reporting_week": "2026-W43",
            "work_completed": "Gate 1 and Gate 2 cameras installed, testing feed.",
            "progress_percent": 65,
            "status": "WIP",
            "remarks": "On schedule for final inspection.",
        },
        user=excel_setup["user1"],
    )
    assert update_obj.progress_percent == 65

    # 3. Verify changes persisted to the single underlying task record
    action.refresh_from_db()
    assert action.current_progress_percent == 65
    assert action.status == "WIP"
    assert "Gate 1 and Gate 2" in action.last_update

    # 4. Verify reflected in Audit & Action Owners (my_tasks)
    my_tasks = list(get_governance_actions({"view_mode": "my_tasks"}, excel_setup["user1"]))
    user_task = next(t for t in my_tasks if t.action_id == "ACT-00505")
    assert user_task.current_progress_percent == 65
    assert user_task.status == "WIP"

    # 5. Verify synchronized with Governance and Action (all tasks view for management)
    gov_tasks = list(get_governance_actions({}, excel_setup["admin"]))
    gov_task = next(t for t in gov_tasks if t.action_id == "ACT-00505")
    assert gov_task.current_progress_percent == 65
    assert gov_task.status == "WIP"
    assert gov_task.last_update == user_task.last_update


@pytest.mark.django_db
def test_excel_import_with_week_column_maps_reporting_week(excel_setup):
    """
    Verifies that the existing 'Week' column from uploaded Excel is normalized and mapped correctly
    to GovernanceActionPoint.reporting_week and synchronizes an initial GovernanceWeeklyUpdate.
    """
    from apps.governance.models import GovernanceWeeklyUpdate

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Action Tracker"

    headers = [
        "Action ID", "Circle", "Function/Area", "Action Point", "Key Requirement / Deliverable",
        "Responsibility", "Meeting / Raised Date ", "Target Closure Date ", "Status", "Priority",
        "Actual Closure Date", "CBH Update – Action Taken ", "Week"
    ]
    ws.append(headers)

    today = timezone.localdate()
    target = today + timedelta(days=14)

    ws.append([
        "ACT-WK42-01",
        "RJ",
        "SCM",
        "Procure spares for inverter cooling",
        "Deliver 50 fan kits",
        "mgupta",
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        "WIP",
        "High",
        "",
        "Supplier purchase order confirmed",
        "Week-42",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is True, res.get("errors")
    assert res["created_count"] == 1

    action = GovernanceActionPoint.objects.get(action_id="ACT-WK42-01")
    assert action.reporting_week == "Week 42"
    assert action.status == "WIP"

    # Verify initial weekly update record was created
    w_upd = GovernanceWeeklyUpdate.objects.filter(action_point=action, reporting_week="Week 42").first()
    assert w_upd is not None
    assert w_upd.status == "WIP"
    assert "Supplier purchase order" in w_upd.work_completed


@pytest.mark.django_db
def test_multi_week_excel_import_and_metrics_filtering(excel_setup):
    """
    Verifies:
    1. Multi-week Excel data upload with Week 42 and Week 43.
    2. Available weeks dynamically populated from database records.
    3. Metrics and actions scoped strictly to the selected week.
    4. Unrecorded weeks show clean zero state.
    5. Underlying task status is never mutated when switching weeks.
    6. Double-counting is prevented when tasks have multiple updates.
    """
    from apps.governance.selectors.governance_selectors import (
        get_governance_metadata,
        get_governance_summary_metrics,
        get_governance_actions,
    )
    from apps.governance.services.governance_weekly_update_service import GovernanceWeeklyUpdateService

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Action Tracker"

    headers = [
        "Action ID", "Circle", "Function/Area", "Action Point", "Key Requirement / Deliverable",
        "Owner", "Raised Date", "Target Closure Date", "Status", "Priority",
        "Actual Closure Date", "Last Update", "Week"
    ]
    ws.append(headers)

    today = timezone.localdate()
    target_future = today + timedelta(days=10)
    target_past = today - timedelta(days=5)

    # Week 42 task (Overdue WIP)
    ws.append([
        "ACT-WK42-001",
        "RJ",
        "Operations",
        "Inverter inspection Week 42",
        "Inspect all stations",
        "mgupta",
        today.strftime("%Y-%m-%d"),
        target_past.strftime("%Y-%m-%d"),
        "WIP",
        "High",
        "",
        "Initial check completed",
        "Week 42",
    ])

    # Week 43 task (Open on-track)
    ws.append([
        "ACT-WK43-002",
        "MH",
        "Finance",
        "Audit compliance Week 43",
        "Complete reconciliation",
        "admin_xl",
        today.strftime("%Y-%m-%d"),
        target_future.strftime("%Y-%m-%d"),
        "Open",
        "Medium",
        "",
        "Pending bank statements",
        "Week 43",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is True
    assert res["created_count"] == 2

    # 1. Available weeks populated dynamically in reverse chronological order
    meta = get_governance_metadata(excel_setup["admin"])
    assert "available_weeks" in meta
    assert "Week 43" in meta["available_weeks"]
    assert "Week 42" in meta["available_weeks"]
    assert meta["available_weeks"].index("Week 43") < meta["available_weeks"].index("Week 42")

    # 2. Week 42 scoped metrics
    metrics_w42 = get_governance_summary_metrics(excel_setup["admin"], reporting_week="Week 42")
    assert metrics_w42["total_actions"] == 1
    assert metrics_w42["wip_count"] == 1
    assert metrics_w42["open_count"] == 0
    assert metrics_w42["overdue_count"] == 1
    assert metrics_w42["reporting_week"] == "Week 42"
    assert len(metrics_w42["circle_distribution"]) == 1
    assert metrics_w42["circle_distribution"][0]["circle__code"] == "RJ"

    actions_w42 = list(get_governance_actions({"reporting_week": "Week 42"}, excel_setup["admin"]))
    assert len(actions_w42) == 1
    assert actions_w42[0].action_id == "ACT-WK42-001"

    # 3. Week 43 scoped metrics
    metrics_w43 = get_governance_summary_metrics(excel_setup["admin"], reporting_week="Week 43")
    assert metrics_w43["total_actions"] == 1
    assert metrics_w43["open_count"] == 1
    assert metrics_w43["wip_count"] == 0
    assert metrics_w43["overdue_count"] == 0
    assert metrics_w43["reporting_week"] == "Week 43"
    assert len(metrics_w43["circle_distribution"]) == 1
    assert metrics_w43["circle_distribution"][0]["circle__code"] == "MH"

    actions_w43 = list(get_governance_actions({"reporting_week": "Week 43"}, excel_setup["admin"]))
    assert len(actions_w43) == 1
    assert actions_w43[0].action_id == "ACT-WK43-002"

    # 4. Empty week returns clean zero values
    metrics_empty = get_governance_summary_metrics(excel_setup["admin"], reporting_week="Week 99")
    assert metrics_empty["total_actions"] == 0
    assert metrics_empty["open_count"] == 0
    assert metrics_empty["wip_count"] == 0
    assert metrics_empty["closed_count"] == 0
    assert metrics_empty["overdue_count"] == 0
    assert metrics_empty["reporting_week"] == "Week 99"
    assert metrics_empty["circle_distribution"] == []
    assert metrics_empty["owner_distribution"] == []
    assert metrics_empty["target_timeline"] == []
    assert metrics_empty["urgent_items"] == []

    # 5. Switching reporting week does NOT mutate underlying task record
    act_obj = GovernanceActionPoint.objects.get(action_id="ACT-WK42-001")
    assert act_obj.status == "WIP"
    assert act_obj.reporting_week == "Week 42"

    # 6. Multi-week progress updates on the same task do not duplicate in aggregations
    GovernanceWeeklyUpdateService.submit_weekly_update(
        action_point=act_obj,
        data={
            "reporting_week": "Week 43",
            "work_completed": "Follow-up testing in Week 43",
            "progress_percent": 80,
            "status": "WIP",
        },
        user=excel_setup["user1"],
    )

    # Act now has updates in Week 42 and Week 43
    # Check Week 43 metrics: act_obj should be included once, plus ACT-WK43-002 => total 2
    metrics_w43_updated = get_governance_summary_metrics(excel_setup["admin"], reporting_week="Week 43")
    assert metrics_w43_updated["total_actions"] == 2
    actions_w43_list = list(get_governance_actions({"reporting_week": "Week 43"}, excel_setup["admin"]))
    assert len(actions_w43_list) == 2
    assert len(set(a.id for a in actions_w43_list)) == 2


@pytest.mark.django_db
def test_excel_import_fallback_reporting_week_from_raised_date(excel_setup):
    """
    Verifies that when an Excel workbook lacks an explicit 'Week' column (e.g. template 3),
    the importer automatically derives reporting_week from the raised_date's ISO week,
    preventing records from having empty reporting_week strings.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Open Points Tracker"

    # Row 3 headers WITHOUT Week column (exactly matching Governance_Action_Tracking_Template (3).xlsx)
    headers = [
        "S.No", "Circle", "Function/Area", "Action Point", "Key Requirement / Deliverable",
        "Owner", "Raised Date", "Target Closure Date", "Current Target", "Revision Count",
        "Status", "Actual Closure Date", "Ageing (Days)", "Ageing Bucket", "Priority",
        "Escalate to Business Head?", "Weeks in Review", "Last Update"
    ]
    ws.append(["Title"])
    ws.append(["Instructions"])
    ws.append(headers)

    # Add a row raised on 2026-10-12 (ISO Week 42)
    ws.append([
        "ACT-FALLBACK-01", "RJ", "Operations", "Fallback Week derivation test", "Deliverable 1",
        "mgupta", "2026-10-12", "2026-10-26", "2026-10-26", 0,
        "Open", "", 0, "0-7 Days", "High", "N", 0, "Initial status"
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is True
    assert res["created_count"] == 1

    act = GovernanceActionPoint.objects.get(action_id="ACT-FALLBACK-01")
    assert act.reporting_week == "Week 42"
    assert act.weekly_updates.filter(reporting_week="Week 42").exists()


@pytest.mark.django_db
def test_api_actions_and_metrics_week_filtering_and_empty_state(excel_setup):
    """
    Verifies that GET /api/v1/governance/actions/ and /metrics/ strictly filter by reporting_week.
    Week 42 returns Week 42 records, Week 43 returns Week 43 records, and unpopulated week
    returns clean empty state without leaking records from other weeks.
    """
    from datetime import date
    from rest_framework.test import APIClient
    client = APIClient()
    client.force_authenticate(user=excel_setup["admin"])

    # Create distinct records for Week 42 and Week 43
    GovernanceActionPoint.objects.create(
        action_id="ACT-W42-TEST",
        circle=excel_setup["circle1"],
        action_point="Action for Week 42 only",
        owner=excel_setup["user1"],
        assigned_by=excel_setup["admin"],
        raised_date=date(2026, 10, 12),
        target_closure_date=date(2026, 10, 26),
        current_target_date=date(2026, 10, 26),
        reporting_week="Week 42",
        status=GovernanceActionPoint.Status.OPEN,
        created_by=excel_setup["admin"],
    )
    GovernanceActionPoint.objects.create(
        action_id="ACT-W43-TEST",
        circle=excel_setup["circle2"],
        action_point="Action for Week 43 only",
        owner=excel_setup["user1"],
        assigned_by=excel_setup["admin"],
        raised_date=date(2026, 10, 19),
        target_closure_date=date(2026, 11, 2),
        current_target_date=date(2026, 11, 2),
        reporting_week="Week 43",
        status=GovernanceActionPoint.Status.WIP,
        created_by=excel_setup["admin"],
    )

    # 1. Query Week 42 via API
    res_actions_42 = client.get("/api/v1/governance/actions/?reporting_week=Week+42", HTTP_HOST="localhost")
    assert res_actions_42.status_code == 200
    action_ids_42 = [item["action_id"] for item in res_actions_42.data["results"]]
    assert "ACT-W42-TEST" in action_ids_42
    assert "ACT-W43-TEST" not in action_ids_42

    res_metrics_42 = client.get("/api/v1/governance/metrics/?reporting_week=Week+42", HTTP_HOST="localhost")
    assert res_metrics_42.status_code == 200
    assert res_metrics_42.data["metrics"]["reporting_week"] == "Week 42"
    assert res_metrics_42.data["metrics"]["total_actions"] == 1
    assert res_metrics_42.data["metrics"]["open_count"] == 1
    assert res_metrics_42.data["metrics"]["wip_count"] == 0

    # 2. Query Week 43 via API
    res_actions_43 = client.get("/api/v1/governance/actions/?reporting_week=Week+43", HTTP_HOST="localhost")
    assert res_actions_43.status_code == 200
    action_ids_43 = [item["action_id"] for item in res_actions_43.data["results"]]
    assert "ACT-W43-TEST" in action_ids_43
    assert "ACT-W42-TEST" not in action_ids_43

    res_metrics_43 = client.get("/api/v1/governance/metrics/?reporting_week=Week+43", HTTP_HOST="localhost")
    assert res_metrics_43.status_code == 200
    assert res_metrics_43.data["metrics"]["reporting_week"] == "Week 43"
    assert res_metrics_43.data["metrics"]["total_actions"] == 1
    assert res_metrics_43.data["metrics"]["open_count"] == 0
    assert res_metrics_43.data["metrics"]["wip_count"] == 1

    # 3. Query Week 99 (Empty state)
    res_actions_99 = client.get("/api/v1/governance/actions/?reporting_week=Week+99", HTTP_HOST="localhost")
    assert res_actions_99.status_code == 200
    assert res_actions_99.data["count"] == 0
    assert len(res_actions_99.data["results"]) == 0

    res_metrics_99 = client.get("/api/v1/governance/metrics/?reporting_week=Week+99", HTTP_HOST="localhost")
    assert res_metrics_99.status_code == 200
    assert res_metrics_99.data["metrics"]["total_actions"] == 0
    assert res_metrics_99.data["metrics"]["reporting_week"] == "Week 99"
    assert res_metrics_99.data["metrics"]["circle_distribution"] == []

    # 4. Switch back to Week 42 to verify idempotency
    res_actions_42_repeat = client.get("/api/v1/governance/actions/?reporting_week=Week+42", HTTP_HOST="localhost")
    assert res_actions_42_repeat.status_code == 200
    assert [item["action_id"] for item in res_actions_42_repeat.data["results"]] == ["ACT-W42-TEST"]


@pytest.mark.django_db
def test_template_and_export_include_week_column(excel_setup):
    """
    Verifies that generated template and export workbook contain the 'Week' and 'Bucket' column headers.
    """
    # Template
    template_stream = GovernanceExcelService.generate_template()
    wb_template = openpyxl.load_workbook(template_stream)
    ws_template = wb_template["Open Points Tracker"]
    headers_template = [cell.value for cell in ws_template[3] if cell.value]
    assert "Week" in headers_template
    assert "Bucket" in headers_template
    assert headers_template[-1] == "Bucket"

    # Export
    export_stream = GovernanceExcelService.export_governance_workbook(excel_setup["admin"])
    wb_export = openpyxl.load_workbook(export_stream)
    ws_export = wb_export["Open Points Tracker"]
    headers_export = [cell.value for cell in ws_export[3] if cell.value]
    assert "Week" in headers_export
    assert "Bucket" in headers_export
    assert headers_export[-1] == "Bucket"


@pytest.mark.django_db
def test_excel_bucket_import_and_backward_compatibility(excel_setup):
    """
    1. Upload with Bucket column correctly maps and saves Bucket value.
    2. Upload with older Excel file without Bucket column still succeeds gracefully.
    3. Re-upload or edit does not erase existing Bucket value.
    4. Ageing Bucket and Bucket remain separate and never conflict.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Open Points Tracker"

    headers = [
        "Action ID", "Circle", "Function/Area", "Action Point", "Key Requirement / Deliverable",
        "Owner", "Raised Date", "Target Closure Date", "Status", "Priority",
        "Actual Closure Date", "Last Update", "Week", "Bucket"
    ]
    ws.append(headers)

    today = timezone.localdate()
    target = today + timedelta(days=14)

    # Row 1 with Bucket "Critical Infrastructure"
    ws.append([
        "",
        "RJ",
        "Operations",
        "Replace degraded battery banks",
        "DC backup online",
        "mgupta",
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        "Open",
        "High",
        "",
        "Vendor selected",
        "Week 42",
        "Critical Infrastructure",
    ])

    # Row 2 with Bucket "Statutory Compliance"
    ws.append([
        "",
        "RJ",
        "Compliance/Legal",
        "Renew environmental NOC certification",
        "Approval document filed",
        "mgupta",
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        "Open",
        "Medium",
        "",
        "Application submitted",
        "Week 42",
        "Statutory Compliance",
    ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = GovernanceExcelService.import_excel(stream, excel_setup["admin"])
    assert res["success"] is True, res.get("errors")
    assert res["created_count"] == 2

    act1 = GovernanceActionPoint.objects.get(action_point="Replace degraded battery banks")
    assert act1.bucket == "Critical Infrastructure"
    assert act1.ageing_bucket == "0-7 Days"  # Dynamic derived metric is separate

    act2 = GovernanceActionPoint.objects.get(action_point="Renew environmental NOC certification")
    assert act2.bucket == "Statutory Compliance"

    # Backward compatibility: Upload older Excel file WITHOUT Bucket column
    wb_old = openpyxl.Workbook()
    ws_old = wb_old.active
    ws_old.title = "Action Tracker"
    old_headers = [
        "Action ID", "Circle", "Function/Area", "Action Point", "Key Requirement / Deliverable",
        "Owner", "Raised Date", "Target Closure Date", "Status", "Priority",
        "Actual Closure Date", "Last Update"
    ]
    ws_old.append(old_headers)
    ws_old.append([
        act1.action_id,
        "RJ",
        "Operations",
        "Replace degraded battery banks",
        "DC backup online",
        "mgupta",
        today.strftime("%Y-%m-%d"),
        target.strftime("%Y-%m-%d"),
        "WIP",
        "High",
        "",
        "Batteries delivered to site",
    ])

    stream_old = io.BytesIO()
    wb_old.save(stream_old)
    stream_old.seek(0)

    res_old = GovernanceExcelService.import_excel(stream_old, excel_setup["admin"])
    assert res_old["success"] is True, res_old.get("errors")
    assert res_old["updated_count"] == 1

    act1.refresh_from_db()
    assert act1.status == "WIP"
    assert act1.last_update == "Batteries delivered to site"
    # Preserves saved Bucket value even when older Excel without Bucket column is uploaded
    assert act1.bucket == "Critical Infrastructure"




